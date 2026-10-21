# -*- coding: utf-8 -*-
"""يولّد ملفات Excel الستة من data.json (واجهة RTL، رؤوس مجمدة، فلاتر، تلوين شرطي)."""
import os, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from common import DATA_JSON, OUT, load_json

D = load_json(DATA_JSON)
FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN, YELLOW, ORANGE, RED, GRAY = 'C6EFCE', 'FFEB9C', 'F8CBAD', 'FFC7CE', 'E7E6E6'


def bullets(lst, mark='• '):
    lst = [x for x in lst if x]
    if not lst:
        return ''
    if len(lst) == 1:
        return lst[0]
    return '\n'.join(mark + x for x in lst)


def new_ws(wb, title, headers, widths, first=False, freeze_cols=1, hidden_last=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = '1F4E5F'
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=11)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = ws.cell(row=2, column=freeze_cols + 1)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'
    if hidden_last:
        ws.column_dimensions[get_column_letter(len(headers))].hidden = True
    return ws


def add_row(ws, vals, link_cols=()):
    r = ws.max_row + 1
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(vertical='top', wrap_text=True, horizontal='right')
        c.border = BORDER
        if i in link_cols and isinstance(v, str) and v.startswith('http') and ' ' not in v and len(v) < 2000:
            c.hyperlink = v
            c.font = Font(name=FONT, size=10, color='0563C1', underline='single')
    return r


def cond(ws, col, rules, last=None):
    """rules: [(نص, لون)] تلوين شرطي بالبحث عن النص في الخلية."""
    last = last or max(ws.max_row, 2)
    L = get_column_letter(col)
    rng = f'{L}2:{L}{last + 200}'
    for text, color in rules:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",{L}2))'],
                                                       fill=PatternFill('solid', bgColor=color, fgColor=color), stopIfTrue=True))


MAT_RULES = [('غير مقيّم', GRAY), ('ضعيفة', RED), ('أولية', ORANGE), ('تحتاج', YELLOW), ('متوسطة', YELLOW), ('جاهزة', GREEN), ('متقدمة', GREEN)]
SEV_RULES = [('حرجة', RED), ('عالية', ORANGE), ('متوسطة', YELLOW), ('منخفضة', GREEN)]
VER_RULES = [('يحتاج تحققًا', YELLOW), ('غير موثق', RED), ('منقول', ORANGE), ('موثق نصًّا', GREEN), ('موثق جزئيًا', 'DDEBF7')]
PRI_RULES = [('عالية', RED), ('متوسطة', YELLOW), ('منخفضة', GREEN)]


def save(wb, name):
    p = os.path.join(OUT, name)
    wb.save(p)
    print('saved', p)


def case_line(c):
    u = f" — {c['urls'][0]}" if c.get('urls') else ''
    return f"{c['name']} [{c['level']}]{u}"


# ------------------------------------------------------------------ 1. خريطة الرسالة
def map_workbook():
    wb = Workbook()
    ws = new_ws(wb, 'دليل', ['البند', 'الشرح'], [32, 110], first=True)
    n_items = sum(len(s['items']) for c in D['chapters'] for s in c['sections'])
    guide = [
        ('الموضوع', D['meta']['title']),
        ('تاريخ التوليد', D['meta']['generated']),
        ('مصدر المحتوى', D['meta']['note']),
        ('طريقة القراءة', 'لكل فصل ورقة؛ كل صف مطلب (أو تمهيد أو بند) بعناوين الهيكل المعتمد حرفيًّا. الأعمدة: الأولوية، درجة النضج، ما هو موجود، الإشكالات والثغرات، ما يمكن إضافته، الأحكام القضائية المرتبطة، المصادر الفقهية المتحققة، المواد القانونية، المصادر والروابط، حالة التوثيق.'),
        ('درجة النضج', 'أخضر: جاهزة أو متقدمة · أصفر: تحتاج أو متوسطة · برتقالي: أولية · أحمر: ضعيفة · رمادي: غير مقيّم. في الفصلين الثالث والرابع الدرجة من جدول الموازن والناسب؛ وفي غيرهما تقدير الباني المشتق من وصف التقرير (يُذكر في عمود «أساس درجة النضج»).'),
        ('رموز الأولوية', '★ مطلب مفصل (نحو 700 كلمة) · ○ مطلب مختصر (نحو 200 كلمة) بحسب قرار الباحث 19/10/2026.'),
        ('الوسوم', 'المحتوى يحمل وسوم التقارير: [نص] [منقول] [تحليل] [استنتاج] [رأي] [يحتاج تحققًا]. ما لا رابط له ذُكر «بلا رابط [يحتاج تحققًا]».'),
        ('عدد البنود', f'{n_items} بندًا في سبع مجموعات (الفصول الخمسة، والخاتمة، والملاحق).'),
        ('تقرير المنهجي (13)', 'موجود وقد أُدرجت ورقتا «فقه الأقليات والقواعد» و«رسائل مشابهة».' if D['meta']['has_report13'] else 'لم يوجد بعد؛ شغّل run_all.py ثانية بعد وروده لتمتلئ ورقتا «فقه الأقليات والقواعد» و«رسائل مشابهة».'),
        ('تنبيه', 'المعرّفات التقنية محفوظة في عمود أخير مخفي في كل ورقة؛ لا تظهر رموز الإشكاليات في الخانات المقروءة.'),
    ]
    for g in guide:
        add_row(ws, list(g))
    ws.auto_filter.ref = None

    headers = ['الفصل', 'المبحث', 'المطلب', 'الأولوية', 'درجة النضج', 'أساس درجة النضج', 'ما هو موجود',
               'الإشكالات والثغرات', 'ما يمكن إضافته', 'الأحكام القضائية المرتبطة', 'المصادر الفقهية المتحققة',
               'المواد القانونية', 'المصادر والروابط', 'حالة التوثيق', 'المعرّف التقني']
    widths = [22, 26, 44, 9, 13, 24, 48, 64, 56, 48, 44, 34, 48, 28, 10]
    sheet_names = {'ch1': 'الفصل 1', 'ch2': 'الفصل 2', 'ch3': 'الفصل 3', 'ch4': 'الفصل 4', 'ch5': 'الفصل 5'}
    groups = [(['ch1'], 'الفصل 1'), (['ch2'], 'الفصل 2'), (['ch3'], 'الفصل 3'), (['ch4'], 'الفصل 4'), (['ch5'], 'الفصل 5'), (['conc', 'app'], 'الخاتمة والملاحق')]
    for ids, title in groups:
        ws = new_ws(wb, title, headers, widths, freeze_cols=3, hidden_last=True)
        for c in D['chapters']:
            if c['id'] not in ids:
                continue
            for s in c['sections']:
                for it in s['items']:
                    cases = [case_line(x) for x in it['cases']]
                    news = [f"إشكالية جديدة مقترحة: {n['title']} ({n['star']})" for n in it['new_issues']]
                    add_row(ws, [
                        c['title'], s['title'] or '—', it['title'] + ((' — ' + ' | '.join(it['outline_details'])) if it['outline_details'] else ''),
                        it['star'] or '—',
                        it['maturity'] + (f" ({it['maturity_score']})" if it['maturity_score'] else ''),
                        it['maturity_basis'] + (f" · الحجم: {it['size']}" if it['size'] else ''),
                        it['exists'] + (f"\nالشاهد في المخطط: {it['witness']}" if it['witness'] else ''),
                        bullets(it['issues'] + ([f"ما ينقص (بحسب الجدول): {it['missing']}"] if it['missing'] else [])),
                        bullets(it['additions'] + news),
                        bullets(cases),
                        bullets([f"{f['t']}" + (f" — {f['u']}" if f['u'] else ' — بلا رابط [يحتاج تحققًا]') for f in it['fiqh']]),
                        bullets(it['laws']),
                        bullets([f"{f['t']}" + (f" — {f['u']}" if f['u'] else ' — بلا رابط [يحتاج تحققًا]') for f in it['sources']]),
                        it['verify'], it['id']])
        cond(ws, 5, MAT_RULES)
        cond(ws, 14, VER_RULES)
    # درجات الفصل الرابع
    ws = new_ws(wb, 'درجات الفصل 4', ['المطلب', 'الواقع', 'الفقه', 'القانون', 'القضاء', 'المقارنة', 'الحكم الإجمالي', 'ما ينقص'], [46, 9, 9, 9, 9, 9, 24, 70], freeze_cols=1)
    titles = {i['id']: i['title'] for c in D['chapters'] for s in c['sections'] for i in s['items']}
    for sc in D['nasib_scores']:
        it = next(i for c in D['chapters'] for s in c['sections'] for i in s['items'] if i['id'] == sc['code'])
        add_row(ws, [it['title'], sc['real'], sc['fiqh'], sc['law'], sc['case'], sc['comp'], sc['verdict'], it['missing']])
    from openpyxl.formatting.rule import ColorScaleRule
    ws.conditional_formatting.add('B2:F40', ColorScaleRule(start_type='num', start_value=0, start_color='F8CBAD', mid_type='num', mid_value=2.5, mid_color='FFEB9C', end_type='num', end_value=5, end_color='C6EFCE'))
    cond(ws, 7, MAT_RULES)
    # مصفوفة الفصل الخامس
    ws = new_ws(wb, 'مصفوفة الفصل 5', ['#', 'التحدي', 'مصدره (إشكاليات)', 'المقترح', 'الجهة', 'الجدوى القانونية', 'الدليل'], [5, 38, 44, 48, 18, 44, 40])
    for r in D['matrix']:
        add_row(ws, r)
    ws = new_ws(wb, 'بنية الملحق الأول', ['القسم', 'المحتوى', 'الجواز القانوني'], [30, 60, 60])
    for r in D['annex1']:
        add_row(ws, r)
    # فقه الأقليات والقواعد
    ws = new_ws(wb, 'فقه الأقليات والقواعد', ['القسم', 'العنوان أو القاعدة', 'الصيغة أو الحالة', 'المرجع أو ما تحقق', 'تطبيقها على الإشكاليات'], [20, 36, 56, 56, 56])
    for r in D['minority_refs']:
        add_row(ws, ['مرجع في فقه الأقليات (المؤسس)', r[0], r[1], r[2], ''])
    if D['rules']:
        for r in D['rules']:
            add_row(ws, ['قاعدة فقهية متحققة (المنهجي)', r[0], r[1], r[2], r[3]])
    else:
        add_row(ws, ['قواعد الترجيح', 'ينتظر تقرير المنهجي (13)', '', '', ''])
    ws = new_ws(wb, 'رسائل مشابهة', ['#', 'الرسالة', 'الدرجة والجامعة والسنة والإشراف', 'الرابط', 'الاطلاع'], [5, 56, 56, 56, 30], freeze_cols=2)
    if D['theses']:
        for r in D['theses']:
            add_row(ws, r, link_cols=(4,))
    else:
        add_row(ws, ['', 'ينتظر تقرير المنهجي (13)', '', '', ''])
    save(wb, 'خريطة_الرسالة.xlsx')


# ------------------------------------------------------------------ 2. الإشكاليات الجديدة
def new_issues_workbook():
    wb = Workbook()
    h = ['الرقم', 'الأولوية', 'عنوان الإشكالية', 'السيناريو', 'الموضع المقترح في الرسالة', 'المطالب المعتمدة المرتبطة',
         'المادة القانونية', 'الفقه', 'الشاهد', 'الأهمية', 'الحل المقترح', 'حالة التوثيق', 'الروابط']
    ws = new_ws(wb, 'الإشكاليات الجديدة', h, [7, 9, 40, 48, 40, 44, 56, 56, 34, 34, 34, 20, 56], first=True, freeze_cols=3)
    titles = {i['id']: i['title'] for c in D['chapters'] for s in c['sections'] for i in s['items']}
    for n in D['new_issues']:
        rel = [titles.get(i, i) for i in n['issue_ids']]
        add_row(ws, [n['n'], n['star'], n['title'], n['scenario'], n['place'], bullets(rel), n['article'], n['fiqh'], n['witness'],
                     n['importance'], n['proposal'], n['verify'], '\n'.join(n['urls'])])
    cond(ws, 12, VER_RULES)
    ws = new_ws(wb, 'اجتماعية من القاضي', ['#', 'إشكالية محتملة في المجتمع الأرجنتيني (تحليل لا حكم)'], [6, 140])
    for i, s in enumerate(D['social_issues'], 1):
        add_row(ws, [i, s])
    save(wb, 'الإشكاليات_الجديدة.xlsx')


# ------------------------------------------------------------------ 3. الأحكام القضائية
def cases_workbook():
    wb = Workbook()
    titles = {i['id']: i['title'] for c in D['chapters'] for s in c['sections'] for i in s['items']}
    h = ['الحكم (المحكمة والقضية)', 'النوع', 'التاريخ', 'رقم Fallos', 'الموضوع', 'الواقعة والمبدأ', 'المطالب المرتبطة', 'مستوى التحقق',
         'ما يحتاج تحققًا', 'الرابط', 'المعرّف التقني']
    ws = new_ws(wb, 'أحكام إضافية (القاضي)', h, [52, 10, 16, 16, 34, 70, 48, 34, 34, 60, 8], first=True, hidden_last=True)
    for c in D['cases_extra']:
        add_row(ws, [c['name'], c['kind'], c['date'], c['fallos'], c['category'], c['summary'],
                     bullets([titles.get(i, i) for i in c['issue_ids']]), c['level'], c['pending'] or ('نعم' if c['need_check'] else ''),
                     '\n'.join(c['urls']) or 'بلا رابط [يحتاج تحققًا]', c['id']], link_cols=(10,))
    cond(ws, 8, [('نص كامل', GREEN), ('رسمي', 'DDEBF7'), ('ثانوي', YELLOW), ('غير', ORANGE)])
    h2 = ['الحكم', 'المحكمة', 'التاريخ', 'الموضوع', 'مستوى التحقق (ملخص/نص كامل)', 'المطالب المرتبطة', 'الرابط', 'المعرّف التقني']
    ws = new_ws(wb, 'أحكام السجل المشترك', h2, [56, 30, 20, 44, 40, 48, 60, 8], hidden_last=True)
    for c in D['cases_old']:
        add_row(ws, [c['name'], c['court'], c['date'], c['topic'], c['level'], bullets(c['issues']), c['url'] or 'بلا رابط [يحتاج تحققًا]', c['id']], link_cols=(7,))
    cond(ws, 5, [('نص كامل', GREEN), ('ملخص', YELLOW), ('صحفي', ORANGE), ('منقول', ORANGE)])
    ws = new_ws(wb, 'فجوات قضائية موثقة', ['#', 'موضوع لم يوجد فيه حكم أرجنتيني (يلزم فحص SAIJ وelDial وCIJ وFallos مباشرة)'], [6, 140])
    for i, g in enumerate(D['case_gaps'], 1):
        add_row(ws, [i, g])
    save(wb, 'الأحكام_القضائية.xlsx')


# ------------------------------------------------------------------ 4. سجل الأخطاء
def errors_workbook():
    wb = Workbook()
    h = ['الرقم', 'الموضع', 'الوصف', 'الخطورة', 'الوكيل', 'التوصية', 'الحالة']
    ws = new_ws(wb, 'السجل', h, [6, 46, 80, 11, 14, 56, 22], first=True, freeze_cols=2)
    order = {'حرجة': 0, 'عالية': 1, 'متوسطة': 2, 'منخفضة': 3}
    errs = sorted(D['errors'], key=lambda e: (order[e['severity']], e['agent_no']))
    for i, e in enumerate(errs, 1):
        add_row(ws, [i, e['position'], e['desc'], e['severity'], e['agent'], e['rec'], e['status']])
    cond(ws, 4, SEV_RULES)
    ws2 = new_ws(wb, 'ملخص بحسب الوكيل', ['الوكيل', 'حرجة', 'عالية', 'متوسطة', 'منخفضة', 'المجموع'], [16, 9, 9, 9, 9, 10])
    agents = []
    for e in D['errors']:
        if e['agent'] not in agents:
            agents.append(e['agent'])
    for a in agents:
        row = [sum(1 for e in D['errors'] if e['agent'] == a and e['severity'] == s) for s in ['حرجة', 'عالية', 'متوسطة', 'منخفضة']]
        add_row(ws2, [a] + row + [sum(row)])
    r = add_row(ws2, ['المجموع'] + [sum(1 for e in D['errors'] if e['severity'] == s) for s in ['حرجة', 'عالية', 'متوسطة', 'منخفضة']] + [len(D['errors'])])
    for c in ws2[r]:
        c.font = Font(name=FONT, bold=True, size=10)
    ws3 = new_ws(wb, 'يحتاج تحققًا (مجمّع)', ['الوكيل', 'البند الذي يحتاج تحققًا'], [16, 140])
    for n in D['needs_check']:
        for x in n['items']:
            add_row(ws3, [n['agent'], x])
    save(wb, 'سجل_الأخطاء_والثغرات.xlsx')


# ------------------------------------------------------------------ 5. المصادر
def sources_workbook():
    wb = Workbook()
    h = ['المصدر أو سياقه', 'النوع', 'الرابط', 'مستوى التحقق', 'الوكيل أو التقرير', 'المطلب المرتبط']
    ws = new_ws(wb, 'المصادر والروابط', h, [64, 28, 70, 40, 24, 50], first=True)
    for s in D['sources']:
        add_row(ws, [s['name'], s['kind'], s['url'] or 'بلا رابط [يحتاج تحققًا]', s['level'], '، '.join(s['agents']), bullets(s.get('related', []))], link_cols=(3,))
    cond(ws, 4, VER_RULES + [('مطابق', GREEN), ('نص مقروء', GREEN), ('ملخص', YELLOW), ('ثانوي', YELLOW)])
    cond(ws, 3, [('بلا رابط', RED)])
    save(wb, 'المصادر_والروابط.xlsx')


# ------------------------------------------------------------------ 6. خطة الرفع
def plan_workbook():
    wb = Workbook()
    h = ['المرحلة', 'المستوى', 'المهمة', 'الأولوية', 'المدة', 'معيار القبول', 'الوكيل المسؤول']
    ws = new_ws(wb, 'مراحل المحرر', h, [12, 24, 64, 11, 18, 60, 14], first=True, freeze_cols=1)
    for s in D['plan_stages']:
        add_row(ws, [s['stage'], s['level'], s['task'], s['priority'], s['duration'], s['criterion'], s['agent']])
    cond(ws, 4, PRI_RULES)
    r = add_row(ws, ['', '', 'الجدول الزمني الإجمالي بحسب المحرر: نحو 22–26 أسبوعًا (المرحلتان 1–4 في الأسابيع 1–4، و5–8 في الأسابيع 5–20، والبقية بعدها) [رأي المحرر].', '', '', '', ''])
    h2 = ['المهمة', 'المرحلة المقابلة', 'الأولوية', 'المدة (من المرحلة)', 'معيار القبول (من المرحلة)', 'الوكيل المسؤول', 'التقرير المصدر']
    ws = new_ws(wb, 'مهام من التقارير', h2, [72, 14, 11, 18, 56, 14, 14], freeze_cols=1)
    for t in sorted(D['plan_tasks'], key=lambda t: ({'عالية': 0, 'متوسطة': 1, 'منخفضة': 2}[t['priority']], t['stage'])):
        add_row(ws, [t['task'], t['stage'], t['priority'], t['duration'], t['criterion'], t['agent'], t['report']])
    cond(ws, 3, PRI_RULES)
    save(wb, 'خطة_الرفع.xlsx')


def main():
    map_workbook(); new_issues_workbook(); cases_workbook(); errors_workbook(); sources_workbook(); plan_workbook()


if __name__ == '__main__':
    main()
